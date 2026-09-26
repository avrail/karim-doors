"""Extract historical reference cases without checking business prices into Git."""

import json
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook


CASES = [
    # Sheet, code(s), customer, width, height, fire minutes, wood rows,
    # other material rows, fees (manufacturing, transport, installation),
    # wood waste percentage, final calculated price row, quote rounding.
    (0, ["RED-D01-FD30-1050"], "RED", 1050, 2200, 30, range(4, 10), [14, 15, 17, 18, 20], (22, 23, 24), 20, 27, 0),
    (1, ["RED-D02-1050", "RED-D03-1050", "RED-D29-1050"], "RED", 1050, 2200, 0, range(4, 13), [17, 18, 19, 21, 22], (25, 26, 27), 10, 30, 0),
    (2, ["RED-D04-950"], "RED", 950, 2200, 0, range(4, 13), [16, 17, 18, 20, 21], (24, 25, 26), 10, 29, 0),
    (3, ["RED-D02-850", "RED-D04-850"], "RED", 850, 2200, 0, range(4, 13), [16, 17, 18, 20, 21], (24, 25, 26), 10, 29, 0),
    (5, ["HA-D01-FD30-1100"], "HA", 1100, 2200, 30, range(4, 12), [15, 16, 17, 18, 20, 21, 22, 24], (26, 27, 28), 10, 31, 2),
    (6, ["HA-D04"], "HA", 970, 2200, 0, range(4, 15), [18, 19, 20, 21, 22, 24, 25], (28, 29, 30), 15, 33, 2),
    (7, ["HA-D03-920"], "HA", 920, 2200, 0, range(4, 15), [18, 19, 20, 21, 22, 24, 25], (28, 29, 30), 15, 33, 2),
    (8, ["HA-D02-870"], "HA", 870, 2200, 0, range(4, 15), [18, 19, 20, 21, 22, 24, 25], (28, 29, 30), 15, 33, 2),
]

ENGLISH = {
    "قائم الحلق": "Frame vertical", "رأس الحلق": "Frame head",
    "قائم البر": "Architrave vertical", "رأس البر": "Architrave head",
    "قائم القشاط": "Oak casing vertical", "رأس القشاط": "Oak casing head",
    "قائم شاسيه": "Leaf stile", "رأس شاسيه": "Leaf rail",
    "السؤاسات": "Cross rails", "قائم علفة كملة حلق": "Frame completion vertical",
    "راس علفة كملة حلق": "Frame completion head",
}


def number(value, sheet, row):
    if not isinstance(value, (int, float)):
        raise ValueError(f"{sheet.title}!L{row} has no cached numeric value")
    return Decimal(str(value))


def as_text(value):
    return str(value or "").strip()


def component(sheet, row, waste, wood, group):
    cost = number(sheet.cell(row, 12).value, sheet, row)
    label = as_text(sheet.cell(row, 3).value) or as_text(sheet.cell(row, 2).value) or as_text(sheet.cell(row, 1).value)
    if wood:
        qty = Decimal(str(sheet.cell(row, 5).value))
        dims = [Decimal(str(sheet.cell(row, col).value)) for col in (6, 7, 8)]
        measurement = qty * dims[0] * dims[1] * dims[2]
        kind = "OAK" if "أرو" in group or "ارو" in group else "MOUSKI"
        unit = "CubicMeter"
        english = ENGLISH.get(label, label)
    else:
        qty = sheet.cell(row, 5).value
        measurement = Decimal(str(qty)) if isinstance(qty, (int, float)) and qty > 0 else Decimal(1)
        if "Halspan" in label:
            kind, unit, english = "FIRECORE", "Piece", "Halspan Optima 30 fire core"
        elif "ابلاكاج" in label:
            kind, unit, english = "PLYWOOD6", "Piece", "Plywood 6 mm"
        elif "MDF18" in label:
            kind, unit, english = "MDF18", "Piece", "MDF 18 mm"
        elif "MDF" in label:
            kind, unit, english = "MDF6", "Piece", "MDF 6 mm"
        elif "قشرة" in label:
            kind, unit, english = "VENEER", "SquareMeter", "Oak veneer"
        elif "فوم" in label:
            kind, unit, english = "FOAM", "Piece", "Foam"
        elif "gasket" in label.lower():
            kind, unit, english = "GASKET", "LinearMeter", "Weather gasket"
        elif "دهانات" in label:
            kind, unit, english = "PAINT", "SquareMeter", "Paint materials"
        elif "كارتون" in label:
            kind, unit, english = "PACKAGING", "Kilogram", "Packaging cardboard"
        else:
            raise ValueError(f"Unclassified component {sheet.title}!{row}: {label}")
    price = cost / measurement
    return {
        "code": f"ROW-{row}", "nameEn": english, "nameAr": label,
        "kind": kind, "unit": unit, "measurement": str(measurement),
        "unitPrice": str(price.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)),
        "wastePercentage": str(waste if wood else 0),
        "sourceRow": row,
    }


def main():
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "docs/source-data/Break Dowen -10-2022.xlsx"
    target = Path(sys.argv[2]) if len(sys.argv) > 2 else source.with_suffix(".reference-cases.json")
    wb = load_workbook(source, read_only=True, data_only=True)
    cases = []
    for index, codes, customer, width, height, fire, wood_rows, other_rows, fee_rows, waste, final_row, rounding in CASES:
        sheet = wb.worksheets[index]
        rows = []
        group = ""
        for row in wood_rows:
            group_text = as_text(sheet.cell(row, 1).value) + " " + as_text(sheet.cell(row, 2).value)
            if "الموسكي" in group_text or "الأرو" in group_text or "الارو" in group_text:
                group = group_text
            rows.append(component(sheet, row, waste, True, group))
        rows.extend(
            component(sheet, row, waste, False, "")
            for row in other_rows
            if isinstance(sheet.cell(row, 12).value, (int, float))
        )
        expected = number(sheet.cell(final_row, 12).value, sheet, final_row)
        for code in codes:
            cases.append({
                "code": code, "customer": customer, "sourceSheet": sheet.title,
                "effectiveDate": "2021-05-10" if customer == "RED" else "2022-01-22",
                "widthMm": width, "heightMm": height, "fireRatingMinutes": fire,
                "manufacturingCost": str(number(sheet.cell(fee_rows[0], 12).value, sheet, fee_rows[0])),
                "transportCost": str(number(sheet.cell(fee_rows[1], 12).value, sheet, fee_rows[1])),
                "installationCost": str(number(sheet.cell(fee_rows[2], 12).value, sheet, fee_rows[2])),
                "administrativePercentage": "17.5", "profitPercentage": "25",
                "quoteRoundingDigits": rounding, "expectedCalculatedPrice": str(expected),
                "components": rows,
            })
    target.write_text(json.dumps({"source": source.name, "cases": cases}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Extracted {len(cases)} reference cases from {len(CASES)} worksheets to {target}")


if __name__ == "__main__":
    main()
